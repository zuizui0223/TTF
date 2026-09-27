#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import shutil
import tempfile
import zipfile
from pathlib import Path


FIGURE_FILES = {
    1: "Figure1_taxonomic_vs_geographic_specialization.png",
    2: "Figure2_anthropogenic_resource_expansion.png",
    3: "Figure3_host_contribution_architecture.png",
    4: "Figure4_within_family_specialization_hierarchy.png",
    5: "Figure5_independent_climate_filtering.png",
}

FORBIDDEN_IDENTITY_TOKENS = (
    "zuizui0223",
    "TTF repository",
    ".github/workflows/",
    "Repository provenance",
)


def split_submission_sections(markdown: str) -> dict[str, str]:
    refs = markdown.find("## References")
    data = markdown.find("## Data and Code Availability")
    figures = markdown.find("## Figure legends")
    if min(refs, data, figures) < 0:
        raise ValueError("missing References / Data and Code Availability / Figure legends")
    if not refs < data < figures:
        raise ValueError(
            "GEB blinded-main order must be References -> Data and Code -> Figure legends"
        )
    return {
        "main_through_data": markdown[:figures].rstrip() + "\n",
        "figure_legends": markdown[figures:].strip(),
    }


def extract_figure_legends(block: str) -> dict[int, str]:
    lines = block.splitlines()
    if not lines or not lines[0].startswith("## Figure legends"):
        raise ValueError("figure legend block must start with ## Figure legends")
    legends: dict[int, str] = {}
    current: list[str] = []
    number: int | None = None

    def flush() -> None:
        nonlocal current, number
        if number is not None:
            legends[number] = " ".join(part.strip() for part in current if part.strip())
        current = []
        number = None

    for raw in lines[1:]:
        line = raw.strip()
        match = re.match(r"^\*\*Figure\s+(\d+)\.", line)
        if match:
            flush()
            number = int(match.group(1))
            current = [line]
        elif number is not None and line:
            current.append(line)
    flush()

    if sorted(legends) != [1, 2, 3, 4, 5]:
        raise ValueError(f"expected figure legends 1-5, got {sorted(legends)}")
    return legends


def keyword_items(markdown: str) -> list[str]:
    match = re.search(r"^\*\*Keywords:\*\*\s*(.+)$", markdown, re.MULTILINE)
    if not match:
        raise ValueError("keywords line not found")
    if ";" in match.group(1):
        raise ValueError("GEB keywords must be comma-separated")
    return [item.strip() for item in match.group(1).split(",") if item.strip()]


def _set_run_font(run, name: str, size_pt: float) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = name
    run.font.size = Pt(size_pt)
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    rfonts = run._element.rPr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|\x60.+?\x60)")


def add_inline_runs(paragraph, text: str, *, size_pt: float = 12.0) -> None:
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            _set_run_font(run, "Times New Roman", size_pt)
        token = match.group(0)
        if token.startswith("**") and token.endswith("**"):
            value = token[2:-2]
            run = paragraph.add_run(value)
            run.bold = True
            _set_run_font(run, "Times New Roman", size_pt)
        elif token.startswith("*") and token.endswith("*"):
            value = token[1:-1]
            run = paragraph.add_run(value)
            run.italic = True
            _set_run_font(run, "Times New Roman", size_pt)
        else:
            value = token[1:-1]
            run = paragraph.add_run(value)
            _set_run_font(run, "Courier New", max(9.0, size_pt - 1.0))
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        _set_run_font(run, "Times New Roman", size_pt)


def _configure_styles(document) -> None:
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Inches, Pt

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(6)

    title = document.styles["Title"]
    title.font.name = "Times New Roman"
    title.font.size = Pt(14)
    title.font.bold = True
    title.paragraph_format.space_after = Pt(12)

    h1 = document.styles["Heading 1"]
    h1.font.name = "Times New Roman"
    h1.font.size = Pt(12)
    h1.font.bold = True
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    h1.paragraph_format.keep_with_next = True

    h2 = document.styles["Heading 2"]
    h2.font.name = "Times New Roman"
    h2.font.size = Pt(12)
    h2.font.bold = True
    h2.font.italic = True
    h2.paragraph_format.space_before = Pt(9)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.keep_with_next = True

    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)


def _enable_continuous_line_numbering(section) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    sect_pr = section._sectPr
    for child in list(sect_pr):
        if child.tag == qn("w:lnNumType"):
            sect_pr.remove(child)
    line_numbers = OxmlElement("w:lnNumType")
    line_numbers.set(qn("w:countBy"), "1")
    line_numbers.set(qn("w:start"), "1")
    line_numbers.set(qn("w:restart"), "continuous")
    sect_pr.append(line_numbers)


def _add_page_number_footer(section) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    _set_run_font(run, "Times New Roman", 10)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def _render_markdown_lines(document, text: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    in_code = False
    current_heading = ""
    for raw in text.splitlines():
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("~~~") or stripped.startswith(chr(96) * 3):
            in_code = not in_code
            continue
        if not stripped or stripped == "---":
            continue

        if in_code:
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(line)
            _set_run_font(run, "Courier New", 9)
            continue

        if line.startswith("# "):
            p = document.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline_runs(p, line[2:].strip(), size_pt=14)
            continue

        if line.startswith("## "):
            heading = line[3:].strip()
            current_heading = heading
            document.add_paragraph(heading, style="Heading 1")
            continue

        if line.startswith("### "):
            heading = line[4:].strip()
            document.add_paragraph(heading, style="Heading 2")
            continue

        if line.startswith("> "):
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.right_indent = Inches(0.2)
            add_inline_runs(p, line[2:].strip())
            for run in p.runs:
                run.italic = True
            continue

        if re.match(r"^\d+\.\s+", line):
            p = document.add_paragraph(style="List Number")
            add_inline_runs(p, re.sub(r"^\d+\.\s+", "", line))
            continue

        if line.startswith("- "):
            if current_heading.startswith("References"):
                p = document.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.first_line_indent = Inches(-0.25)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(3)
                add_inline_runs(p, line[2:].strip(), size_pt=11)
            else:
                p = document.add_paragraph(style="List Bullet")
                add_inline_runs(p, line[2:].strip())
            continue

        p = document.add_paragraph()
        add_inline_runs(p, line)


def _set_image_alt_text(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description[:900])


def _add_figures(document, legends: dict[int, str], figures_dir: Path) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    document.add_paragraph("Figure legends", style="Heading 1")
    for number in range(1, 6):
        image = figures_dir / FIGURE_FILES[number]
        if not image.exists():
            raise FileNotFoundError(image)

        caption = document.add_paragraph()
        caption.paragraph_format.page_break_before = True
        caption.paragraph_format.keep_with_next = True
        caption.paragraph_format.line_spacing = 1.0
        caption.paragraph_format.space_after = Pt(6)
        add_inline_runs(caption, legends[number], size_pt=10)

        picture_p = document.add_paragraph()
        picture_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_p.paragraph_format.keep_together = True
        run = picture_p.add_run()
        shape = run.add_picture(str(image), width=Inches(6.15))
        _set_image_alt_text(
            shape,
            f"Figure {number}",
            re.sub(r"\*+", "", legends[number]),
        )


def _scrub_docx_metadata(path: Path) -> None:
    tmp = Path(tempfile.mkstemp(suffix=".docx")[1])
    try:
        with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(
            tmp, "w", compression=zipfile.ZIP_DEFLATED
        ) as zout:
            for item in zin.infolist():
                if item.filename == "docProps/custom.xml":
                    continue
                data = zin.read(item.filename)
                if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                    data = re.sub(
                        rb'\s+w:rsid[A-Za-z0-9]*="[^"]*"',
                        b"",
                        data,
                    )
                if item.filename == "docProps/core.xml":
                    data = re.sub(
                        rb"<dc:creator>.*?</dc:creator>",
                        b"<dc:creator></dc:creator>",
                        data,
                        flags=re.DOTALL,
                    )
                    data = re.sub(
                        rb"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>",
                        b"<cp:lastModifiedBy></cp:lastModifiedBy>",
                        data,
                        flags=re.DOTALL,
                    )
                if item.filename == "docProps/app.xml":
                    data = re.sub(
                        rb"<Company>.*?</Company>",
                        b"<Company></Company>",
                        data,
                        flags=re.DOTALL,
                    )
                zout.writestr(item, data)
        shutil.move(str(tmp), str(path))
    finally:
        if tmp.exists():
            tmp.unlink(missing_ok=True)


def validate_review_docx(path: Path) -> dict[str, object]:
    with zipfile.ZipFile(path, "r") as zf:
        names = set(zf.namelist())
        document_xml = zf.read("word/document.xml")
        core_xml = zf.read("docProps/core.xml")
        media = sorted(name for name in names if name.startswith("word/media/"))
        combined = b"\n".join(
            zf.read(name)
            for name in names
            if name.endswith(".xml") and zf.getinfo(name).file_size < 5_000_000
        )

    if b"w:lnNumType" not in document_xml:
        raise RuntimeError("continuous line numbering is absent")
    if len(media) != 5:
        raise RuntimeError(f"expected exactly five embedded figures, found {len(media)}")
    if re.search(rb"<dc:creator>\s*[^<]+", core_xml):
        raise RuntimeError("creator metadata is not blank")
    if re.search(rb"<cp:lastModifiedBy>\s*[^<]+", core_xml):
        raise RuntimeError("lastModifiedBy metadata is not blank")

    lower = combined.lower()
    for token in FORBIDDEN_IDENTITY_TOKENS:
        if token.lower().encode("utf-8") in lower:
            raise RuntimeError(f"review DOCX contains forbidden identity token: {token}")

    return {
        "schema": "ttf_butterfly_specialization_geb_review_docx_receipt_v0.1",
        "status": "GENERATED_ANONYMIZED_REVIEW_DOCX",
        "docx": str(path),
        "embedded_figures": len(media),
        "continuous_line_numbering": True,
        "author_metadata_blank": True,
        "forbidden_identity_tokens_absent": True,
    }


def render_review_docx(markdown_path: Path, figures_dir: Path, output: Path) -> dict[str, object]:
    from docx import Document

    markdown = markdown_path.read_text(encoding="utf-8")
    for token in FORBIDDEN_IDENTITY_TOKENS:
        if token.lower() in markdown.lower():
            raise ValueError(f"blinded Markdown contains forbidden identity token: {token}")

    keywords = keyword_items(markdown)
    if not 6 <= len(keywords) <= 10:
        raise ValueError(f"GEB requires 6-10 keywords; got {len(keywords)}")
    if keywords != sorted(keywords, key=str.casefold):
        raise ValueError("GEB keywords are not alphabetized")

    sections = split_submission_sections(markdown)
    legends = extract_figure_legends(sections["figure_legends"])

    document = Document()
    _configure_styles(document)
    section = document.sections[0]
    _enable_continuous_line_numbering(section)
    _add_page_number_footer(section)

    props = document.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = "Butterfly specialization is hierarchical"
    props.subject = "Blinded review manuscript"
    props.comments = ""

    _render_markdown_lines(document, sections["main_through_data"])
    _add_figures(document, legends, figures_dir)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    _scrub_docx_metadata(output)
    return validate_review_docx(output)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--markdown", type=Path, required=True)
    ap.add_argument("--figures-dir", type=Path, required=True)
    ap.add_argument("--output", type=Path, required=True)
    ap.add_argument("--receipt", type=Path, required=True)
    args = ap.parse_args()

    receipt = render_review_docx(args.markdown, args.figures_dir, args.output)
    args.receipt.parent.mkdir(parents=True, exist_ok=True)
    args.receipt.write_text(
        json.dumps(receipt, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(receipt, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
