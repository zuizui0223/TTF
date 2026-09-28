#!/usr/bin/env python3
from __future__ import annotations

import argparse
import shutil
import subprocess
import zipfile
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
MAIN_MD = ROOT / "manuscript/ttf_q_methods_blinded_v0.1.md"
TITLE_MD = ROOT / "manuscript/ttf_q_title_page_template.md"
LICENSE = ROOT / "LICENSE"

TITLE_PLACEHOLDERS = (
    "[AUTHOR",
    "[AFFILIATION",
    "[INSTITUTION",
    "[NAME]",
    "[POSTAL ADDRESS]",
    "[EMAIL]",
    "[COMPLETE FOR TITLE PAGE",
    "[CRediT",
    "[DECLARATION]",
)


def readiness_blockers() -> list[str]:
    blockers: list[str] = []
    if not LICENSE.is_file():
        blockers.append("missing_root_LICENSE")
    if not MAIN_MD.is_file():
        blockers.append("missing_blinded_main_document")
    if not TITLE_MD.is_file():
        blockers.append("missing_title_page")
    elif any(token in TITLE_MD.read_text() for token in TITLE_PLACEHOLDERS):
        blockers.append("unresolved_title_page_placeholders")
    if shutil.which("pandoc") is None:
        blockers.append("pandoc_not_available")
    return blockers


def add_page_field(paragraph) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, value, end):
        run._r.append(node)


def make_reference_docx(path: Path) -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is required for final MEE rendering"
        ) from exc

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_after = Pt(0)

    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.paragraph_format.line_spacing = 2.0

    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:lnNumType"))
    if existing is not None:
        sect_pr.remove(existing)
    line_numbers = OxmlElement("w:lnNumType")
    line_numbers.set(qn("w:countBy"), "1")
    line_numbers.set(qn("w:distance"), "360")
    line_numbers.set(qn("w:restart"), "continuous")
    sect_pr.append(line_numbers)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(footer)

    document.add_paragraph("REFERENCE DOCUMENT FOR TTF-Q SUBMISSION RENDERING")
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def assert_submission_docx(path: Path, *, require_line_numbers: bool) -> None:
    if not path.is_file() or path.stat().st_size == 0:
        raise RuntimeError(f"missing rendered DOCX: {path}")
    with zipfile.ZipFile(path) as zf:
        document_xml = zf.read("word/document.xml").decode("utf-8")
        footer_names = [name for name in zf.namelist() if name.startswith("word/footer")]
        footer_xml = "\n".join(zf.read(name).decode("utf-8") for name in footer_names)
    if require_line_numbers:
        if "w:lnNumType" not in document_xml or 'w:restart="continuous"' not in document_xml:
            raise RuntimeError("continuous line numbering missing from Main Document")
    if " PAGE " not in footer_xml:
        raise RuntimeError("page-number PAGE field missing from footer")


def run_pandoc(source: Path, output: Path, reference: Path) -> None:
    subprocess.run(
        [
            "pandoc",
            str(source),
            "--from=markdown",
            "--to=docx",
            f"--reference-doc={reference}",
            "--standalone",
            f"--output={output}",
        ],
        check=True,
        cwd=ROOT,
    )


def convert_pdf(docx: Path, output_dir: Path) -> Path | None:
    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    if libreoffice is None:
        return None
    subprocess.run(
        [
            libreoffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_dir),
            str(docx),
        ],
        check=True,
        cwd=ROOT,
    )
    pdf = output_dir / f"{docx.stem}.pdf"
    return pdf if pdf.is_file() else None


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()

    blockers = readiness_blockers()
    if blockers:
        raise RuntimeError(
            "TTF-Q final submission rendering is blocked: " + ", ".join(blockers)
        )

    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    reference = output_dir / "ttf_q_reference.docx"
    main_docx = output_dir / "ttf_q_main_document_blinded.docx"
    title_docx = output_dir / "ttf_q_title_page.docx"

    make_reference_docx(reference)
    run_pandoc(MAIN_MD, main_docx, reference)
    run_pandoc(TITLE_MD, title_docx, reference)

    assert_submission_docx(main_docx, require_line_numbers=True)
    assert_submission_docx(title_docx, require_line_numbers=False)

    main_pdf = convert_pdf(main_docx, output_dir)
    title_pdf = convert_pdf(title_docx, output_dir)

    print(
        {
            "status": "RENDERED_TTF_Q_MEE_SUBMISSION_FILES",
            "main_docx": str(main_docx),
            "title_docx": str(title_docx),
            "main_pdf": None if main_pdf is None else str(main_pdf),
            "title_pdf": None if title_pdf is None else str(title_pdf),
        }
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
