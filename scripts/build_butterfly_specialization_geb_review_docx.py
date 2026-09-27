#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
import tempfile
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


FIGURE_PREFIXES = {
    1: "Figure1_taxonomic_vs_geographic_specialization",
    2: "Figure2_anthropogenic_resource_expansion",
    3: "Figure3_host_contribution_architecture",
    4: "Figure4_within_family_specialization_hierarchy",
    5: "Figure5_independent_climate_filtering",
}


def _section(text: str, heading: str, next_headings: tuple[str, ...]) -> str:
    start = text.find(heading)
    if start < 0:
        raise ValueError(f"missing section: {heading}")
    ends = [
        text.find(next_heading, start + len(heading))
        for next_heading in next_headings
    ]
    ends = [value for value in ends if value >= 0]
    end = min(ends) if ends else len(text)
    return text[start:end].strip() + "\n"


def _figure_legends(text: str) -> dict[int, str]:
    block = _section(
        text,
        "## Figure legends",
        ("## Data and Code Availability", "## References"),
    )
    block = block[len("## Figure legends"):].strip()
    matches = list(
        re.finditer(
            r"\*\*Figure\s+(\d+)\.\s+(.*?)\*\*\s*(.*?)(?=\n\n\*\*Figure\s+\d+\.|\Z)",
            block,
            flags=re.DOTALL,
        )
    )
    legends = {}
    for match in matches:
        number = int(match.group(1))
        title = re.sub(r"\s+", " ", match.group(2).strip())
        body = re.sub(r"\s+", " ", match.group(3).strip())
        legends[number] = f"**Figure {number}. {title}** {body}".strip()
    if set(legends) != set(FIGURE_PREFIXES):
        raise ValueError(f"expected five figure legends; got {sorted(legends)}")
    return legends


def build_submission_markdown(source: str, figure_dir: Path) -> str:
    fig_start = source.find("## Figure legends")
    if fig_start < 0:
        raise ValueError("missing Figure legends section")
    body = source[:fig_start].rstrip() + "\n"

    references = _section(source, "## References (working)", ("## Repository provenance",))
    references = references.replace("## References (working)", "## References", 1)

    data_code = _section(
        source,
        "## Data and Code Availability",
        ("## References (working)", "## References"),
    )

    legends = _figure_legends(source)

    chunks = [
        body,
        references,
        data_code,
        "## Figures\n",
    ]
    for number in range(1, 6):
        prefix = FIGURE_PREFIXES[number]
        candidates = sorted(figure_dir.glob(prefix + ".png"))
        if len(candidates) != 1:
            raise ValueError(
                f"expected one PNG for {prefix}; found {len(candidates)}"
            )
        image_path = candidates[0].resolve()
        chunks.extend(
            [
                f"### Figure {number}\n",
                f"![Figure {number}]({image_path.as_posix()}){{width=6.35in}}\n",
                legends[number] + "\n",
            ]
        )
    return "\n".join(chunk.rstrip() for chunk in chunks if chunk.strip()) + "\n"


def _add_line_numbering(document: Document) -> None:
    for section in document.sections:
        sect_pr = section._sectPr
        existing = sect_pr.find(qn("w:lnNumType"))
        if existing is None:
            existing = OxmlElement("w:lnNumType")
            sect_pr.append(existing)
        existing.set(qn("w:countBy"), "1")
        existing.set(qn("w:restart"), "continuous")
        existing.set(qn("w:distance"), "360")


def _set_review_layout(document: Document) -> None:
    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        if style_name in document.styles:
            style = document.styles[style_name]
            style.font.name = "Times New Roman"

    for paragraph in document.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True


def _format_figure_pages(document: Document) -> None:
    first_figure_seen = False
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        has_drawing = bool(paragraph._p.xpath(".//w:drawing"))
        if text == "Figures":
            paragraph.paragraph_format.page_break_before = True
        if re.fullmatch(r"Figure\s+[1-5]", text):
            if first_figure_seen:
                paragraph.paragraph_format.page_break_before = True
            first_figure_seen = True
            paragraph.paragraph_format.keep_with_next = True
        if has_drawing:
            paragraph.paragraph_format.keep_with_next = True
            paragraph.alignment = 1


def _scrub_metadata(document: Document) -> None:
    props = document.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""
    props.keywords = ""
    props.subject = ""
    props.category = ""
    props.content_status = ""


def build_docx(source: Path, figure_dir: Path, output: Path) -> None:
    manuscript = source.read_text(encoding="utf-8")
    submission_md = build_submission_markdown(manuscript, figure_dir)

    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        intermediate = Path(tmp) / "submission.docx"
        pypandoc.convert_text(
            submission_md,
            to="docx",
            format="gfm+raw_attribute",
            outputfile=str(intermediate),
            extra_args=["--standalone"],
        )
        document = Document(intermediate)

    _set_review_layout(document)
    _add_line_numbering(document)
    _format_figure_pages(document)
    _scrub_metadata(document)

    document.save(output)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--manuscript", type=Path, required=True)
    ap.add_argument("--figure-dir", type=Path, required=True)
    ap.add_argument("--output", type=Path, required=True)
    args = ap.parse_args()

    build_docx(args.manuscript, args.figure_dir, args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
