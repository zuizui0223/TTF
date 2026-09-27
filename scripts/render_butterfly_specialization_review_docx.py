#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


FIGURE_FILES = {
    1: "Figure1_taxonomic_vs_geographic_specialization.png",
    2: "Figure2_anthropogenic_resource_expansion.png",
    3: "Figure3_host_contribution_architecture.png",
    4: "Figure4_within_family_specialization_hierarchy.png",
    5: "Figure5_independent_climate_filtering.png",
}

INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|\x60[^\x60]+\x60)")


def set_font(run, *, name="Times New Roman", size=12, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline(paragraph, text: str, *, size=12):
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_font(run, size=size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, size=size, bold=True)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, size=size, italic=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, name="Courier New", size=max(9, size - 1))
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_font(run, size=size)


def set_line_numbering(section):
    sect_pr = section._sectPr
    for node in list(sect_pr):
        if node.tag == qn("w:lnNumType"):
            sect_pr.remove(node)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:distance"), "360")
    ln.set(qn("w:restart"), "continuous")
    sect_pr.append(ln)


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_font(run, size=10)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def configure_document(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    set_line_numbering(section)
    add_page_number(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size in (("Title", 14), ("Heading 1", 14), ("Heading 2", 12)):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True

    doc.styles["Heading 1"].paragraph_format.keep_with_next = True
    doc.styles["Heading 2"].paragraph_format.keep_with_next = True

    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""
    props.subject = ""
    props.category = ""
    props.keywords = ""


def add_reference(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(4)
    add_inline(p, text, size=11)


def add_figure(doc: Document, number: int, caption_text: str, figures_dir: Path):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    add_inline(p, caption_text, size=10)

    image = figures_dir / FIGURE_FILES[number]
    if not image.exists():
        raise FileNotFoundError(image)
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.space_after = Pt(0)
    run = pic.add_run()
    run.add_picture(str(image), width=Inches(6.05))


def render(markdown: str, figures_dir: Path) -> Document:
    doc = Document()
    configure_document(doc)

    current_section = ""
    in_code = False
    figure_count = 0
    figure_heading_seen = False

    lines = markdown.splitlines()
    for raw in lines:
        line = raw.rstrip()

        if line.startswith(chr(96) * 3):
            in_code = not in_code
            continue

        if in_code:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(line)
            set_font(run, name="Courier New", size=9)
            continue

        if not line:
            continue
        if line.strip() == "---":
            continue

        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, line[2:].strip(), size=14)
            continue

        if line.startswith("## "):
            current_section = line[3:].strip()
            if current_section == "Figure legends":
                doc.add_page_break()
                figure_heading_seen = True
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, current_section, size=14)
            continue

        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, line[4:].strip(), size=12)
            continue

        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.right_indent = Inches(0.2)
            p.paragraph_format.line_spacing = 1.25
            add_inline(p, line[2:].strip(), size=11)
            continue

        if current_section == "Figure legends" and line.startswith("**Figure "):
            m = re.match(r"\*\*Figure\s+(\d+)\.", line)
            if not m:
                raise ValueError(f"could not parse figure number: {line}")
            number = int(m.group(1))
            if figure_count and number != figure_count + 1:
                raise ValueError("figure legend numbering drift")
            if figure_count:
                doc.add_page_break()
            add_figure(doc, number, line, figures_dir)
            figure_count = number
            continue

        if current_section == "References" and line.startswith("- "):
            add_reference(doc, line[2:].strip())
            continue

        if re.match(r"^\d+\.\s+", line):
            body = re.sub(r"^\d+\.\s+", "", line)
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.line_spacing = 1.5
            add_inline(p, body, size=12)
            continue

        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing = 1.5
            add_inline(p, line[2:].strip(), size=12)
            continue

        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(6)
        add_inline(p, line, size=12)

    if figure_count != 5:
        raise RuntimeError(f"expected five embedded figures, got {figure_count}")
    if not figure_heading_seen:
        raise RuntimeError("Figure legends section not found")
    return doc


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", type=Path, required=True)
    ap.add_argument("--figures-dir", type=Path, required=True)
    ap.add_argument("--output", type=Path, required=True)
    args = ap.parse_args()

    text = args.input.read_text(encoding="utf-8")
    required_order = [
        "## References",
        "## Data and Code Availability",
        "## Figure legends",
    ]
    positions = [text.index(marker) for marker in required_order]
    if positions != sorted(positions):
        raise RuntimeError("GEB tail section order drift")

    doc = render(text, args.figures_dir)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
